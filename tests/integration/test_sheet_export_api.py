"""Integration coverage for exporting corrected live database values."""

import os
import shutil
import subprocess
from collections.abc import Iterator
from datetime import date
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from app.database import get_db
from app.exports.database_adapter import build_billing_source
from app.main import app
from app.models import Extraction, Sheet


@pytest.fixture
def db_session() -> Iterator[Session]:
    engine = create_engine(os.environ["DATABASE_URL"])
    connection = engine.connect()
    transaction = connection.begin()
    session = Session(bind=connection, join_transaction_mode="create_savepoint")
    try:
        yield session
    finally:
        session.close()
        if transaction.is_active:
            transaction.rollback()
        connection.close()
        engine.dispose()


@pytest.fixture
def client(db_session: Session) -> Iterator[TestClient]:
    app.dependency_overrides[get_db] = lambda: db_session
    with TestClient(app) as test_client:
        yield test_client
    app.dependency_overrides.clear()


def seed_sheet(db: Session, vehicle: str = "FAKE-001") -> Sheet:
    sheet = Sheet(
        vehicle=vehicle,
        branch="Test branch",
        date=date(2026, 8, 1),
        image_path="/data/uploads/test-sheet.jpg",
        status="review",
    )
    db.add(sheet)
    db.flush()
    return sheet


def seed_row(
    db: Session,
    sheet: Sheet,
    *,
    row_number: int,
    values: dict[str, str],
) -> dict[str, Extraction]:
    extractions: dict[str, Extraction] = {}
    for field_name, value in values.items():
        extraction = Extraction(
            sheet_id=sheet.id,
            row_number=row_number,
            field_name=field_name,
            image_crop_ref=f"/data/crops/row-{row_number}-{field_name}.jpg",
            raw_ocr_value=value,
            confidence=0.9,
            rule_flag=field_name == "start_km",
            rule_flag_reason=("Needs correction" if field_name == "start_km" else None),
        )
        db.add(extraction)
        extractions[field_name] = extraction
    db.flush()
    return extractions


def recalculate_with_libreoffice(source: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    assert soffice, "LibreOffice is required to validate cached Excel formula results"
    output_dir.mkdir()
    profile = output_dir / "profile"
    subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(source),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return output_dir / source.name


def test_export_endpoint_uses_saved_correction_and_recalculates_without_errors(
    client: TestClient,
    db_session: Session,
    tmp_path: Path,
) -> None:
    sheet = seed_sheet(db_session)
    fields = seed_row(
        db_session,
        sheet,
        row_number=1,
        values={
            "ds_no": "1",
            "date": "01/08/2026",
            "start_time": "8.00",
            "start_km": "23,695",
            "close_time": "20:00",
            "close_km": "23,796",
            "total_km": "101",
            "total_time": "12 H",
            "toll": "35.50",
            "parking": "40",
            "journey_details": "Office to airport",
        },
    )

    correction = client.post(
        f"/review/fields/{fields['start_km'].id}/correction",
        json={"final_value": "23,700"},
    )
    assert correction.status_code == 200

    response = client.get(f"/sheets/{sheet.id}/export")

    assert response.status_code == 200
    assert response.headers["content-type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "attachment" in response.headers["content-disposition"]
    assert response.content.startswith(b"PK")

    output = tmp_path / "downloaded.xlsx"
    output.write_bytes(response.content)
    formulas = openpyxl.load_workbook(output, data_only=False)
    worksheet = formulas["FAKE-001"]
    assert worksheet["F4"].value == 23700
    assert worksheet["F4"].value != 23695
    assert worksheet["G4"].value == 23796
    assert worksheet["H4"].value == "=G4-F4"
    assert worksheet["O4"].value == pytest.approx(75.5)
    formulas.close()

    recalculated = recalculate_with_libreoffice(output, tmp_path / "recalculated")
    values = openpyxl.load_workbook(recalculated, data_only=True)
    calculated = values["FAKE-001"]
    assert calculated["H4"].value == 96
    assert not [
        cell.coordinate
        for row in calculated.iter_rows()
        for cell in row
        if cell.data_type == "e"
    ]
    values.close()


def test_database_adapter_groups_and_orders_extractions_by_row_number(
    db_session: Session,
) -> None:
    sheet = seed_sheet(db_session)
    common = {
        "start_time": "08:00",
        "start_km": "100",
        "close_time": "09:00",
        "close_km": "110",
        "total_time": "1h",
    }
    seed_row(
        db_session,
        sheet,
        row_number=2,
        values={"date": "02/08/2026", **common},
    )
    seed_row(
        db_session,
        sheet,
        row_number=1,
        values={"date": "01/08/2026", **common},
    )

    source = build_billing_source(db_session, sheet.id)

    assert [row["date"] for row in source["rows"]] == [
        date(2026, 8, 1),
        date(2026, 8, 2),
    ]


def test_export_endpoint_returns_404_for_unknown_sheet(client: TestClient) -> None:
    assert client.get("/sheets/999999/export").status_code == 404


def test_export_endpoint_returns_422_when_vehicle_has_no_contract_config(
    client: TestClient, db_session: Session
) -> None:
    sheet = seed_sheet(db_session, vehicle="UNCONFIGURED-VEHICLE")
    seed_row(
        db_session,
        sheet,
        row_number=1,
        values={
            "date": "2026-08-01",
            "start_time": "08:00",
            "start_km": "100",
            "close_time": "09:00",
            "close_km": "110",
            "total_time": "1h",
        },
    )

    response = client.get(f"/sheets/{sheet.id}/export")

    assert response.status_code == 422
    assert "contract config" in response.json()["detail"].lower()
