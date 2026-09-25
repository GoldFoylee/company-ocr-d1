from __future__ import annotations

import json
import shutil
import subprocess
from datetime import time, timedelta
from pathlib import Path

import openpyxl
import pytest

from app.exports.billing_excel import create_monthly_billing_export

FIXTURE_DIR = Path(__file__).parents[1] / "fixtures" / "anonymized"
CONTRACT_RATES = Path(__file__).parents[2] / "config" / "contract_rates.json"
FIXTURE_IDS = [f"sheet_{index:03d}" for index in range(1, 6)]


def _recalculate_with_libreoffice(source: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    assert soffice, "LibreOffice is required to validate cached Excel formula results"
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir = output_dir / "lo-profile"
    subprocess.run(
        [
            soffice,
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--headless",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(output_dir),
            str(source),
        ],
        check=True,
        capture_output=True,
        text=True,
    )
    return output_dir / source.name


def _assert_no_formula_errors(sheet: openpyxl.worksheet.worksheet.Worksheet) -> None:
    assert not [
        cell.coordinate
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == "e"
    ]


def _excel_days(value: time | timedelta | float | int) -> float:
    if isinstance(value, timedelta):
        return value.total_seconds() / 86_400
    if isinstance(value, time):
        return (value.hour * 3600 + value.minute * 60 + value.second) / 86_400
    return float(value)


def test_all_fixtures_export_to_vehicle_tabs_and_recalculate(tmp_path: Path) -> None:
    inputs = [FIXTURE_DIR / f"{fixture_id}.json" for fixture_id in FIXTURE_IDS]
    source_data = {}
    for path in inputs:
        source = json.loads(path.read_text())
        source_data[source["sheet_id"]] = source
    contract_config = json.loads(CONTRACT_RATES.read_text())
    vehicle_keys = [
        contract_config["fixture_sheet_to_vehicle"][fixture_id]
        for fixture_id in FIXTURE_IDS
    ]
    output = tmp_path / "monthly-billing.xlsx"
    create_monthly_billing_export(inputs, output, CONTRACT_RATES)

    formulas = openpyxl.load_workbook(output, data_only=False)
    assert formulas.sheetnames == vehicle_keys
    for fixture_id in FIXTURE_IDS:
        sheet = formulas[contract_config["fixture_sheet_to_vehicle"][fixture_id]]
        assert [sheet.cell(3, column).value for column in range(1, 17)] == [
            "Sl. No.",
            "Veh.No",
            "Duty Date",
            "Days",
            "Modal",
            "St.Km",
            "Cl.Km",
            "Total Kms",
            "Start Time",
            "End Time",
            "Total Time",
            "Time Limit",
            "Extra Hrs",
            "Outstation",
            "Park/Toll",
            "Night",
        ]
        first = source_data[fixture_id]["rows"][0]
        first_row = 3 + int(first["date"][-2:])
        assert sheet[f"D{first_row}"].value == f'=UPPER(TEXT(C{first_row},"dddd"))'
        assert sheet[f"H{first_row}"].value == f"=G{first_row}-F{first_row}"
        assert sheet[f"K{first_row}"].value == f"=J{first_row}-I{first_row}"
        assert "MAX(0" in sheet[f"M{first_row}"].value
        assert sheet[f"N{first_row}"].value is None
        assert sheet[f"P{first_row}"].value is None
        assert sheet["H36"].value == "=SUM(H4:H34)"
        assert sheet["M36"].value == "=SUM(M4:M34)"
        assert sheet["O36"].value == "=SUM(O4:O34)"
        assert sheet["J38"].value == "=H36"
        assert sheet["J43"].value == "=MAX(0,J38-J37)"
        assert sheet["J44"].value == "=M36"
        assert sheet["K45"].value == "=O36"
        assert sheet["L48"].value == "=SUM(L42:L47)"
        assert sheet["A51"].value.startswith("PLACEHOLDER CONTRACT RATES")
        assert sheet["K4"].number_format == "[h]:mm"
        assert sheet.max_row == 51
    formulas.close()

    recalculated_file = _recalculate_with_libreoffice(output, tmp_path / "recalculated")
    values = openpyxl.load_workbook(recalculated_file, data_only=True)
    for fixture_id in FIXTURE_IDS:
        source = source_data[fixture_id]
        vehicle_key = contract_config["fixture_sheet_to_vehicle"][fixture_id]
        calculated = values[vehicle_key]
        first = source["rows"][0]
        first_row = 3 + int(first["date"][-2:])
        assert calculated[f"H{first_row}"].value == first["close_km"] - first["start_km"]
        start_at = calculated[f"I{first_row}"].value
        end_at = calculated[f"J{first_row}"].value
        assert _excel_days(calculated[f"K{first_row}"].value) == pytest.approx(
            (end_at - start_at).total_seconds() / 86_400
        )

        expected_km = sum(item["close_km"] - item["start_km"] for item in source["rows"])
        expected_toll = sum(
            (item.get("toll") or 0) + (item.get("parking") or 0) for item in source["rows"]
        )
        assert calculated["H36"].value == expected_km
        assert calculated["O36"].value == expected_toll
        assert calculated["K45"].value == expected_toll
        assert calculated["J39"].value == sum(
            item["close_km"] - item["start_km"] > 0 for item in source["rows"]
        )
        full_day_item = next(
            (item for item in source["rows"] if item.get("total_time") == "1d"), None
        )
        if full_day_item is not None:
            full_day_row = 3 + int(full_day_item["date"][-2:])
            assert _excel_days(calculated[f"K{full_day_row}"].value) == pytest.approx(1.0)
        _assert_no_formula_errors(calculated)
    values.close()


def test_elapsed_time_exceeds_24_hours_and_billing_inputs_are_formula_driven(
    tmp_path: Path,
) -> None:
    source = json.loads((FIXTURE_DIR / "sheet_001.json").read_text())
    source["rows"] = [dict(source["rows"][0])]
    source["rows"][0].update(
        {"close_time": "22:30", "total_time": "38:30", "toll": 35, "parking": 40}
    )
    source["vehicle_no"] = "TEST-UNIT-123"
    config = json.loads(CONTRACT_RATES.read_text())
    fixture_vehicle_key = config["fixture_sheet_to_vehicle"]["sheet_001"]
    vehicle = config["vehicles"].pop(fixture_vehicle_key)
    vehicle["vehicle_no"] = source["vehicle_no"]
    vehicle["time_limit_hours"] = 12
    vehicle["contract_rates"]["total_fix_km"] = 100
    config["vehicles"][source["vehicle_no"]] = vehicle
    vehicle_key = source["vehicle_no"]
    config_path = tmp_path / "contract_rates.json"
    config_path.write_text(json.dumps(config))
    output = tmp_path / "multi-day-billing.xlsx"
    create_monthly_billing_export([source], output, config_path)

    formulas = openpyxl.load_workbook(output, data_only=False)[vehicle_key]
    assert formulas["K4"].value == "=J4-I4"
    assert formulas["M4"].value == '=IF(L4="",0,MAX(0,K4-L4))'
    assert formulas["J43"].value == "=MAX(0,J38-J37)"
    assert formulas["K45"].value == "=O36"

    recalculated_file = _recalculate_with_libreoffice(output, tmp_path / "recalculated")
    values = openpyxl.load_workbook(recalculated_file, data_only=True)[vehicle_key]
    assert _excel_days(values["K4"].value) == pytest.approx(38.5 / 24)
    assert _excel_days(values["M4"].value) == pytest.approx(26.5 / 24)
    assert values["J43"].value == 1
    assert values["K45"].value == 75
    assert values["L45"].value == 75
    _assert_no_formula_errors(values)
